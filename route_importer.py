import hashlib
import json
import os
import re
import tempfile
import unicodedata
from datetime import date, datetime, time, timedelta, timezone
from decimal import Decimal, InvalidOperation
from pathlib import Path
from zoneinfo import ZoneInfo

import firebase_admin
from firebase_admin import credentials, firestore
import gdown
from openpyxl import load_workbook


CONFIG_PATH = "settings/automaticRouteImport"
TIMEZONE = ZoneInfo("America/Sao_Paulo")


def normalized(value):
    text = unicodedata.normalize("NFKD", str(value or ""))
    text = "".join(char for char in text if not unicodedata.combining(char))
    return re.sub(r"\s+", " ", text.strip().lower())


def slug(value):
    return re.sub(r"^-|-$", "", re.sub(r"[^a-z0-9]+", "-", normalized(value)))


def text(value):
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def integer(value):
    if value in (None, ""):
        return 0
    try:
        return round(float(str(value).replace(",", ".")))
    except ValueError as error:
        raise ValueError(f"Quantidade inválida: {value}") from error


def money_cents(value):
    if value in (None, ""):
        return None
    raw = re.sub(r"[^0-9,.-]", "", str(value))
    if "," in raw and "." in raw:
        raw = raw.replace(".", "").replace(",", ".") if raw.rfind(",") > raw.rfind(".") else raw.replace(",", "")
    elif "," in raw:
        raw = raw.replace(",", ".")
    try:
        amount = Decimal(raw)
    except InvalidOperation as error:
        raise ValueError(f"Valor da nota inválido: {value}") from error
    if amount < 0:
        raise ValueError("O valor da nota não pode ser negativo.")
    return int((amount * 100).quantize(Decimal("1")))


def route_date(value):
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    raw = text(value)
    for pattern in ("%d/%m/%Y", "%Y-%m-%d", "%d-%m-%Y"):
        try:
            return datetime.strptime(raw, pattern).date()
        except ValueError:
            pass
    raise ValueError(f"Data inválida na planilha: {raw or 'vazia'}")


def parse_clock(value, fallback="18:00"):
    try:
        return datetime.strptime(str(value or fallback), "%H:%M").time()
    except ValueError as error:
        raise ValueError("Horário inválido na configuração web.") from error


def initialize_firestore():
    raw = os.environ.get("FIREBASE_SERVICE_ACCOUNT", "").strip()
    if not raw:
        raise RuntimeError("O segredo FIREBASE_SERVICE_ACCOUNT não foi configurado no GitHub.")
    firebase_admin.initialize_app(credentials.Certificate(json.loads(raw)))
    return firestore.client()


def should_run(config, now):
    force = os.environ.get("FORCE_RUN", "false").lower() == "true"
    last_run = config.get("lastRunAt")
    requested = config.get("runRequestedAt")
    manual = requested is not None and (last_run is None or requested > last_run)
    target = parse_clock(config.get("executionTime"))
    schedule_key = f"{now.date().isoformat()}_{target.strftime('%H:%M')}"
    scheduled = (
        config.get("enabled") is True
        and now.isoweekday() in [int(day) for day in config.get("days", [])]
        and now.time() >= target
        and config.get("lastScheduleKey") != schedule_key
    )
    return force or manual or scheduled, schedule_key


def download_excel(folder_url, expected_name):
    destination = Path(tempfile.mkdtemp(prefix="vendify-routes-"))
    downloaded = gdown.download_folder(
        url=folder_url,
        output=str(destination),
        quiet=False,
        use_cookies=False,
        remaining_ok=True,
    ) or []
    files = [Path(path) for path in downloaded if str(path).lower().endswith(".xlsx")]
    if expected_name:
        files = [path for path in files if path.name.lower() == expected_name.lower()]
    if not files:
        raise ValueError("Nenhum arquivo Excel .xlsx foi encontrado na pasta compartilhada.")
    if len(files) > 1 and not expected_name:
        names = ", ".join(sorted(path.name for path in files))
        raise ValueError(f"Há mais de um Excel na pasta ({names}). Informe o nome no painel web.")
    return files[0]


def file_hash(path):
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def read_routes(path):
    workbook = load_workbook(path, data_only=True, read_only=True)
    sheet = workbook["ROTAS"] if "ROTAS" in workbook.sheetnames else workbook[workbook.sheetnames[0]]
    rows = sheet.iter_rows(values_only=True)
    try:
        header_row = next(rows)
    except StopIteration as error:
        raise ValueError("A planilha está vazia.") from error
    headers = {normalized(value): index for index, value in enumerate(header_row)}
    required = [
        "rota", "nome da loja", "endereco", "bairro", "cep", "cidade",
        "estado", "congelados", "refrigerados", "mercearia", "outros",
        "total geral", "data", "motorista", "conferente",
    ]
    missing = [name for name in required if name not in headers]
    if missing:
        raise ValueError("Colunas obrigatórias ausentes: " + ", ".join(missing))

    def cell(row, name, default=None):
        index = headers.get(name)
        return default if index is None or index >= len(row) else row[index]

    grouped = {}
    delivery_dates = set()
    for source_order, row in enumerate(rows, start=1):
        route_name = text(cell(row, "rota"))
        store_name = text(cell(row, "nome da loja"))
        if not route_name and not store_name:
            continue
        if store_name.upper().startswith("EXEMPLO -"):
            continue
        if not route_name or not store_name:
            raise ValueError("Existe uma linha sem Rota ou Nome da Loja.")
        delivery_date = route_date(cell(row, "data"))
        delivery_dates.add(delivery_date)
        driver = text(cell(row, "motorista"))
        checker = text(cell(row, "conferente"))
        helper = text(cell(row, "ajudante"))
        if not driver or not checker:
            raise ValueError(f"A rota {route_name} possui motorista ou conferente vazio.")
        key = (delivery_date, route_name)
        group = grouped.setdefault(key, {
            "date": delivery_date,
            "route_name": route_name,
            "driver": driver,
            "checker": checker,
            "helper": helper,
            "stops": [],
        })
        if normalized(group["driver"]) != normalized(driver) or normalized(group["checker"]) != normalized(checker) or normalized(group["helper"]) != normalized(helper):
            raise ValueError(f"A equipe muda entre as linhas da rota {route_name}.")
        group["stops"].append({
            "order": integer(cell(row, "ordem", source_order)) or source_order,
            "name": store_name,
            "street": text(cell(row, "endereco")),
            "neighborhood": text(cell(row, "bairro")),
            "cep": text(cell(row, "cep")),
            "city": text(cell(row, "cidade")),
            "state": text(cell(row, "estado")),
            "boxes": {
                "frozen": integer(cell(row, "congelados")),
                "chilled": integer(cell(row, "refrigerados")),
                "grocery": integer(cell(row, "mercearia")),
                "bakery": integer(cell(row, "padaria")),
                "other": integer(cell(row, "outros")),
            },
            "invoice": money_cents(cell(row, "valor da nota")),
        })
    if not grouped:
        raise ValueError("A planilha não possui rotas preenchidas.")
    if len(delivery_dates) > 1:
        raise ValueError("A planilha possui mais de uma data. Use um Excel por data de entrega.")
    for group in grouped.values():
        group["stops"].sort(key=lambda stop: stop["order"])
    return list(grouped.values())


def users_index(db):
    result = {}
    duplicates = set()
    for snapshot in db.collection("users").stream():
        data = snapshot.to_dict() or {}
        for candidate in (data.get("name"), data.get("email")):
            key = normalized(candidate)
            if not key:
                continue
            if key in result and result[key][0] != snapshot.id:
                duplicates.add(key)
            result[key] = (snapshot.id, data)
    for key in duplicates:
        result.pop(key, None)
    return result


def find_user(index, description, field, route_name):
    found = index.get(normalized(description))
    if found is None:
        raise ValueError(f"{field} “{description}” da rota {route_name} não foi encontrado ou está duplicado no cadastro.")
    return found


class Writer:
    def __init__(self, db):
        self.db = db
        self.batch = db.batch()
        self.count = 0

    def set(self, reference, data):
        if self.count >= 400:
            self.batch.commit()
            self.batch = self.db.batch()
            self.count = 0
        self.batch.set(reference, data)
        self.count += 1

    def finish(self):
        if self.count:
            self.batch.commit()


def publish(db, routes, source_name, source_hash, loading_time):
    index = users_index(db)
    prepared = []
    for route in routes:
        driver_id, driver = find_user(index, route["driver"], "Motorista", route["route_name"])
        checker_id, checker = find_user(index, route["checker"], "Conferente", route["route_name"])
        helper_id = None
        helper = {}
        if route["helper"]:
            helper_id, helper = find_user(index, route["helper"], "Ajudante", route["route_name"])
        date_id = route["date"].isoformat()
        reference = db.collection("routes").document(f"{date_id}_{slug(route['route_name'])}")
        if reference.get().exists:
            raise ValueError(f"A rota {route['route_name']} de {date_id} já existe. Ela não foi sobrescrita.")
        prepared.append((route, reference, driver_id, driver, helper_id, helper, checker_id, checker))

    writer = Writer(db)
    for route, reference, driver_id, driver, helper_id, helper, checker_id, checker in prepared:
        route_day = route["date"]
        members = [driver_id, checker_id] + ([helper_id] if helper_id else [])
        planned_total = sum(sum(stop["boxes"].values()) for stop in route["stops"])
        writer.set(reference, {
            "date": route_day.isoformat(),
            "routeName": route["route_name"],
            "loadingAt": loading_time,
            "driverId": driver_id,
            "driverName": driver.get("name", route["driver"]),
            "helperId": helper_id,
            "helperName": helper.get("name", "Não informado") if helper_id else "Não informado",
            "checkerId": checker_id,
            "checkerName": checker.get("name", route["checker"]),
            "memberIds": members,
            "plannedTotal": planned_total,
            "sourceFile": source_name,
            "sourceFileHash": source_hash,
            "importMethod": "google-drive-automatic",
            "publishedBy": "automatic-import",
            "publishedAt": firestore.SERVER_TIMESTAMP,
            "detailsExpireAt": datetime.combine(route_day + timedelta(days=180), time(), tzinfo=timezone.utc),
            "routeExpiresAt": datetime.combine(route_day + timedelta(days=365), time(), tzinfo=timezone.utc),
        })
        for index_stop, stop in enumerate(route["stops"], start=1):
            stop_id = f"{index_stop:03d}-{slug(stop['name'])}"
            address = f"{stop['street']}, {stop['neighborhood']}, {stop['city']} - {stop['state']}, {stop['cep']}"
            writer.set(reference.collection("stops").document(stop_id), {
                "order": stop["order"],
                "name": stop["name"],
                "address": address,
                "street": stop["street"],
                "neighborhood": stop["neighborhood"],
                "cep": stop["cep"],
                "city": stop["city"],
                "state": stop["state"],
                "boxes": stop["boxes"],
                "loadingAt": loading_time,
            })
            if stop["invoice"] is not None:
                writer.set(reference.collection("managerData").document(stop_id), {
                    "invoiceValueCents": stop["invoice"],
                    "updatedAt": firestore.SERVER_TIMESTAMP,
                })
    writer.finish()
    return len(prepared), sum(len(item[0]["stops"]) for item in prepared)


def main():
    db = initialize_firestore()
    reference = db.document(CONFIG_PATH)
    config = reference.get().to_dict() or {}
    now = datetime.now(TIMEZONE)
    run, schedule_key = should_run(config, now)
    if not run:
        print("Fora do horário configurado ou automação pausada.")
        return
    try:
        path = download_excel(config.get("folderUrl", ""), config.get("fileName", "").strip())
        digest = file_hash(path)
        if digest == config.get("lastFileHash"):
            reference.set({
                "lastRunAt": firestore.SERVER_TIMESTAMP,
                "lastScheduleKey": schedule_key,
                "lastStatus": "Sem arquivo novo",
                "lastFile": path.name,
                "lastMessage": "O Excel já havia sido importado e não foi duplicado.",
            }, merge=True)
            return
        routes = read_routes(path)
        route_count, stop_count = publish(
            db, routes, path.name, digest, config.get("loadingTime", "06:00")
        )
        reference.set({
            "lastRunAt": firestore.SERVER_TIMESTAMP,
            "lastScheduleKey": schedule_key,
            "lastStatus": "Sucesso",
            "lastFile": path.name,
            "lastFileHash": digest,
            "lastMessage": f"{route_count} rota(s) e {stop_count} loja(s) importadas.",
        }, merge=True)
    except Exception as error:
        reference.set({
            "lastRunAt": firestore.SERVER_TIMESTAMP,
            "lastScheduleKey": schedule_key,
            "lastStatus": "Erro",
            "lastMessage": str(error),
        }, merge=True)
        raise


if __name__ == "__main__":
    main()
